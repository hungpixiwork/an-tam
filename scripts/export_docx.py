import argparse
import os
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

PROJECT_DIR = Path(__file__).resolve().parents[1]
AGENTS_DIR = str(PROJECT_DIR / ".agents")
DEFAULT_OUTPUT = PROJECT_DIR / ".agents" / "submission" / "Ho-so-du-an-An-Tam-v3.docx"

# Màu chủ đạo theo phong cách Light Mode thanh lịch
COLOR_PRIMARY = RGBColor(27, 58, 92)      # Navy đậm #1B3A5C
COLOR_SECONDARY = RGBColor(46, 92, 138)   # Xanh vừa #2E5C8A
COLOR_DARK = RGBColor(34, 34, 34)         # Chữ chính #222222
COLOR_MUTED = RGBColor(100, 110, 120)     # Chữ phụ #646E78

HEX_PRIMARY = "1B3A5C"
HEX_LIGHT_BG = "F4F7F9"
HEX_BORDER = "D0D7DE"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_PRIMARY}"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def format_inline(paragraph, text, base_font="Segoe UI", base_size=10.5, default_color=COLOR_DARK):
    # Regex xử lý bold, italic, code
    # Tạm đơn giản hóa: tách tokens
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) >= 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.name = base_font
            run.font.size = Pt(base_size)
            run.font.color.rgb = default_color
        elif token.startswith("*") and token.endswith("*") and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            run.font.name = base_font
            run.font.size = Pt(base_size)
            run.font.color.rgb = default_color
        elif token.startswith("`") and token.endswith("`") and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(base_size - 0.5)
            run.font.color.rgb = RGBColor(160, 40, 60)
        else:
            clean_text = token.replace("\\[", "[").replace("\\]", "]")
            clean_text = clean_text.replace("[x]", "☑").replace("[ ]", "☐")
            run = paragraph.add_run(clean_text)
            run.font.name = base_font
            run.font.size = Pt(base_size)
            run.font.color.rgb = default_color

def add_callout(doc, text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    
    # Border trái đậm màu primary
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_PRIMARY}"/>
            <w:top w:val="none"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    format_inline(p, text, base_size=10, default_color=RGBColor(40, 50, 60))
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def parse_markdown_to_docx(doc, md_content, skip_h1=False):
    lines = md_content.splitlines()
    i = 0
    in_code_block = False
    code_lines = []
    
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Xử lý code block
        if stripped.startswith("```"):
            if in_code_block:
                # Đóng code block
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                set_cell_background(cell, "F6F8FA")
                set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                code_text = "\n".join(code_lines)
                run = p.add_run(code_text)
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(30, 30, 30)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        # Header 1
        if stripped.startswith("# "):
            if not skip_h1:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(16)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.keep_with_next = True
                run = p.add_run(stripped[2:])
                run.bold = True
                run.font.name = "Segoe UI"
                run.font.size = Pt(18)
                run.font.color.rgb = COLOR_PRIMARY
            i += 1
            continue

        # Header 2
        if stripped.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[3:])
            run.bold = True
            run.font.name = "Segoe UI"
            run.font.size = Pt(14)
            run.font.color.rgb = COLOR_SECONDARY
            i += 1
            continue

        # Header 3
        if stripped.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[4:])
            run.bold = True
            run.font.name = "Segoe UI"
            run.font.size = Pt(12)
            run.font.color.rgb = COLOR_PRIMARY
            i += 1
            continue

        # Header 4
        if stripped.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[5:])
            run.bold = True
            run.font.name = "Segoe UI"
            run.font.size = Pt(11)
            run.font.color.rgb = COLOR_SECONDARY
            i += 1
            continue

        # Blockquote (> )
        if stripped.startswith(">"):
            quote_text = stripped.lstrip(">").strip()
            add_callout(doc, quote_text)
            i += 1
            continue

        # Đường kẻ ngang
        if stripped in ["---", "***", "___"]:
            # Khoảng cách nhẹ
            i += 1
            continue

        # Bảng biểu Markdown (| col | col |)
        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            
            # Xử lý bảng
            if len(table_lines) >= 2:
                header_row = [c.strip() for c in table_lines[0].split("|")[1:-1]]
                data_rows = []
                for tline in table_lines[1:]:
                    cells = [c.strip() for c in tline.split("|")[1:-1]]
                    if all(re.match(r"^:?-+:?$", c) for c in cells if c):
                        continue # Bỏ qua dòng separator
                    data_rows.append(cells)
                
                num_cols = len(header_row)
                tbl = doc.add_table(rows=len(data_rows) + 1, cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_borders(tbl)

                # Header
                for col_idx, text in enumerate(header_row):
                    cell = tbl.cell(0, col_idx)
                    set_cell_background(cell, HEX_PRIMARY)
                    set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    format_inline(p, text, base_font="Segoe UI", base_size=9.5, default_color=RGBColor(255, 255, 255))
                    p.runs[0].bold = True if p.runs else False

                # Data rows
                for r_idx, row in enumerate(data_rows):
                    bg_color = HEX_LIGHT_BG if r_idx % 2 == 1 else "FFFFFF"
                    for col_idx in range(num_cols):
                        text = row[col_idx] if col_idx < len(row) else ""
                        cell = tbl.cell(r_idx + 1, col_idx)
                        if bg_color != "FFFFFF":
                            set_cell_background(cell, bg_color)
                        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(2)
                        p.paragraph_format.space_after = Pt(2)
                        format_inline(p, text, base_font="Segoe UI", base_size=9.5, default_color=COLOR_DARK)
                
                p_spacer = doc.add_paragraph()
                p_spacer.paragraph_format.space_before = Pt(0)
                p_spacer.paragraph_format.space_after = Pt(6)
            continue

        # Bullet list hoặc Checkbox
        if stripped.startswith("- ") or stripped.startswith("* ") or re.match(r"^\d+\.\s", stripped):
            is_ordered = bool(re.match(r"^\d+\.\s", stripped))
            if is_ordered:
                prefix_match = re.match(r"^\d+\.\s*", stripped)
                content = stripped[prefix_match.end():]
                prefix = stripped[:prefix_match.end()]
            else:
                prefix = "• "
                content = stripped[2:]
            
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            
            run_p = p.add_run(prefix)
            run_p.bold = True
            run_p.font.name = "Segoe UI"
            run_p.font.size = Pt(10.5)
            run_p.font.color.rgb = COLOR_SECONDARY
            
            format_inline(p, content, base_font="Segoe UI", base_size=10.5, default_color=COLOR_DARK)
            i += 1
            continue

        # Đoạn văn bản bình thường
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.2
        format_inline(p, stripped, base_font="Segoe UI", base_size=10.5, default_color=COLOR_DARK)
        i += 1

def build_cover_page(doc):
    p_top = doc.add_paragraph()
    p_top.paragraph_format.space_before = Pt(40)
    
    # Subtitle / Category
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(10)
    r_sub = p_sub.add_run("WEBANK FINTECHATHON 2026 — DỰ ÁN DỰ THI")
    r_sub.font.name = "Segoe UI"
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = COLOR_SECONDARY
    r_sub.bold = True
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(12)
    r_title = p_title.add_run("AN TÂM")
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(36)
    r_title.font.color.rgb = COLOR_PRIMARY
    r_title.bold = True
    
    # Tagline
    p_tag = doc.add_paragraph()
    p_tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tag.paragraph_format.space_after = Pt(30)
    r_tag = p_tag.add_run("“Trước mỗi quyết định lớn, biết chắc mình có kham nổi không”")
    r_tag.font.name = "Segoe UI"
    r_tag.font.size = Pt(14)
    r_tag.font.color.rgb = RGBColor(70, 80, 90)
    r_tag.italic = True

    # Info Box
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F4F7F9")
    set_cell_margins(cell, top=160, bottom=160, left=240, right=240)
    set_table_borders(tbl)
    
    p_info = cell.paragraphs[0]
    p_info.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_info.paragraph_format.line_spacing = 1.3
    
    lines = [
        ("Nền tảng:", " Quản lý tài chính cá nhân & Đánh giá rủi ro quyết định cuộc đời"),
        ("Mục tiêu:", " Giúp người trẻ vững vàng tài chính, đối mặt rủi ro nghề nghiệp & AI"),
        ("Đội ngũ phát triển:", " Uyên (Lead/Pitch), Linh (Research), Quỳnh Anh (Design), Hưng (Dev)"),
        ("Ngày lập hồ sơ:", " Tháng 9/2026 — Phiên bản 1.0 Final"),
    ]
    for label, val in lines:
        r1 = p_info.add_run(label)
        r1.bold = True
        r1.font.name = "Segoe UI"
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = COLOR_PRIMARY
        r2 = p_info.add_run(val + "\n")
        r2.font.name = "Segoe UI"
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = COLOR_DARK
        
    doc.add_page_break()

def build_executive_summary(doc):
    p_h1 = doc.add_paragraph()
    p_h1.paragraph_format.space_before = Pt(10)
    p_h1.paragraph_format.space_after = Pt(8)
    r = p_h1.add_run("TỔNG HỢP TÍNH NĂNG CỐT LÕI (BẢN TÓM TẮT DỄ HIỂU)")
    r.bold = True
    r.font.name = "Segoe UI"
    r.font.size = Pt(16)
    r.font.color.rgb = COLOR_PRIMARY

    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_after = Pt(12)
    p_intro.paragraph_format.line_spacing = 1.25
    format_inline(p_intro, "**An Tâm** là nền tảng quản lý tài chính thế hệ mới được thiết kế cho người trẻ Việt Nam trước những bước ngoặt lớn trong sự nghiệp và đời sống. Thay vì chỉ ghi chép thu chi thụ động, An Tâm trả lời trực diện câu hỏi: *“Trước mỗi quyết định lớn, bạn có kham nổi không và cần chuẩn bị điều gì?”*")

    # Mục 1: Thang 5 bậc
    p_h2 = doc.add_paragraph()
    p_h2.paragraph_format.space_before = Pt(10)
    p_h2.paragraph_format.space_after = Pt(6)
    r2 = p_h2.add_run("1. Thang 5 Bậc Khả Năng Chống Chịu Tài Chính")
    r2.bold = True
    r2.font.name = "Segoe UI"
    r2.font.size = Pt(13)
    r2.font.color.rgb = COLOR_SECONDARY

    ladder_data = [
        ("Bậc 1", "Dòng tiền sống còn (Cash Flow)", "Hàng tháng thu có đủ bù chi không? Dư được bao nhiêu?"),
        ("Bậc 2", "Đệm an toàn (Runway & Quỹ khẩn cấp)", "Nếu ngày mai mất việc hoặc đổ bệnh, bạn sống được mấy tháng?"),
        ("Bậc 3", "Năng lực kiếm tiền (Giá trị nghề & AI)", "Nghề của bạn quy ra tiền là bao nhiêu? Nguy cơ bị AI thay thế ra sao?"),
        ("Bậc 4", "Tích lũy & Đầu tư (Stress Test)", "Nếu thị trường sập 30%, bạn có hoảng loạn bán tháo cắt lỗ không?"),
        ("Bậc 5", "Tự do lựa chọn (Life Options)", "Tôi có thể nghỉ việc, đi học thạc sĩ, hay khởi nghiệp chưa?"),
    ]
    
    tbl = doc.add_table(rows=len(ladder_data) + 1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl)
    
    headers = ["Bậc", "Tên bậc & Trọng tâm", "Câu hỏi đời thường giải đáp cho bạn"]
    for c_idx, h_text in enumerate(headers):
        c = tbl.cell(0, c_idx)
        set_cell_background(c, HEX_PRIMARY)
        set_cell_margins(c, 100, 100, 120, 120)
        p = c.paragraphs[0]
        r = p.add_run(h_text)
        r.bold = True
        r.font.name = "Segoe UI"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    for r_idx, (b, name, q) in enumerate(ladder_data):
        bg = HEX_LIGHT_BG if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate([b, name, q]):
            c = tbl.cell(r_idx + 1, c_idx)
            if bg != "FFFFFF":
                set_cell_background(c, bg)
            set_cell_margins(c, 80, 80, 120, 120)
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Segoe UI"
            r.font.size = Pt(9.5)
            r.font.color.rgb = COLOR_DARK
            if c_idx == 0:
                r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Mục 2: 4 Đột phá
    p_h2_2 = doc.add_paragraph()
    p_h2_2.paragraph_format.space_before = Pt(12)
    p_h2_2.paragraph_format.space_after = Pt(6)
    r2_2 = p_h2_2.add_run("2. Bốn Tính Năng Đột Phá (Breakthroughs)")
    r2_2.bold = True
    r2_2.font.name = "Segoe UI"
    r2_2.font.size = Pt(13)
    r2_2.font.color.rgb = COLOR_SECONDARY

    bts = [
        ("• B1 - Lựa chọn cuộc đời (Life Options Engine): ", "Không chỉ tính hưu trí xa vời, tính năng này mô phỏng các quyết định lớn như nghỉ việc chuyển ngành, học cao học, khởi nghiệp. Hệ thống đưa ra 3 kết luận rõ ràng: **Kham nổi ngay / Chưa kham nổi / Cần thêm điều kiện gì** (ví dụ để dành thêm bao nhiêu triệu/tháng trong mấy tháng)."),
        ("• B2 - Giá trị nghề nghiệp & Rủi ro AI (Human Capital + AI Exposure): ", "Tài sản lớn nhất của người trẻ là khả năng kiếm tiền trong 20-30 năm tới. An Tâm định giá toàn bộ thu nhập tương lai (Human Capital) và đối chiếu với dữ liệu rủi ro tự động hóa từ ILO 2025 để người dùng thấy mức độ đe dọa hoặc đòn bẩy AI trong ngành của mình, cùng giá trị quy đổi khi dành thời gian nâng cao kỹ năng."),
        ("• B3 - Thử thách áp lực & Đo tâm lý hoảng loạn (Stress Test + Emotional Gap): ", "Mô phỏng tài chính qua các cuộc khủng hoảng thực tế: Lạm phát phi mã, sập sàn chứng khoán, hoặc làn sóng sa thải công nghệ. Đặc biệt, tính toán **Khoảng cách hành vi (Behavior Gap)** – đo lường số tiền bạn sẽ mất nếu tâm lý yếu và bán tháo ở đáy thay vì kiên định kế hoạch."),
        ("• B4 - Tập dợt đầu tư & Cảnh báo rủi ro dồn cục (Try Before Invest + Concentration Alert): ", "Cho phép đầu tư thử nghiệm bằng tiền ảo với giá thị trường thực. Tự động phát hiện rủi ro kép: nếu làm trong ngành công nghệ mà toàn bộ tiền đầu tư và cổ phiếu thưởng cũng dồn vào công nghệ, khi ngành khủng hoảng bạn sẽ vừa mất việc vừa bốc hơi tài sản."),
    ]
    for title, desc in bts:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.2
        r_t = p.add_run(title)
        r_t.bold = True
        r_t.font.name = "Segoe UI"
        r_t.font.size = Pt(10.5)
        r_t.font.color.rgb = COLOR_PRIMARY
        format_inline(p, desc, base_size=10.5)

    # Mục 3: AI
    p_h2_3 = doc.add_paragraph()
    p_h2_3.paragraph_format.space_before = Pt(12)
    p_h2_3.paragraph_format.space_after = Pt(6)
    r2_3 = p_h2_3.add_run("3. Triết Lý AI: Nói Ít, Tính Chuẩn, Bảo Mật Tuyệt Đối")
    r2_3.bold = True
    r2_3.font.name = "Segoe UI"
    r2_3.font.size = Pt(13)
    r2_3.font.color.rgb = COLOR_SECONDARY

    ai_points = [
        ("AI KHÔNG BAO GIỜ làm toán:", " Mọi phép tính số tiền, runway, lợi nhuận đều do thuật toán mã nguồn (Deterministic Engine) thực hiện với độ chính xác 100%."),
        ("Bảo mật số liệu tuyệt đối (Zero Financial Leak):", " Khi người dùng hỏi chatbot, hệ thống gỡ sạch toàn bộ số tiền, thu nhập, số dư trước khi chuyển cho LLM phân loại ý định."),
        ("Trả lời theo mẫu định sẵn:", " Kết quả trả về cho người dùng được ghép từ dữ liệu đã tính toán vào các mẫu câu chuẩn, loại bỏ hoàn toàn hiện tượng AI bịa số (Hallucination)."),
    ]
    for t, d in ai_points:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(3)
        r_bullet = p.add_run("– ")
        r_bullet.bold = True
        r_t = p.add_run(t)
        r_t.bold = True
        r_t.font.name = "Segoe UI"
        r_t.font.size = Pt(10.5)
        r_t.font.color.rgb = COLOR_PRIMARY
        r_d = p.add_run(d)
        r_d.font.name = "Segoe UI"
        r_d.font.size = Pt(10.5)

    doc.add_page_break()

def main():
    parser = argparse.ArgumentParser(description="Xuất hồ sơ dự án An Tâm từ các tài liệu .agents")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Đường dẫn DOCX đầu ra")
    args = parser.parse_args()
    doc = Document()
    
    # Thiết lập lề trang chuẩn (Normal: 1 inch / 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # 1. Trang bìa
    build_cover_page(doc)

    # 2. Trang Mục lục
    p_toc = doc.add_paragraph()
    p_toc.paragraph_format.space_after = Pt(12)
    r_toc = p_toc.add_run("MỤC LỤC HỒ SƠ DỰ ÁN")
    r_toc.bold = True
    r_toc.font.name = "Segoe UI"
    r_toc.font.size = Pt(16)
    r_toc.font.color.rgb = COLOR_PRIMARY

    files = [
        ("TỔNG HỢP TÍNH NĂNG CỐT LÕI (BẢN TÓM TẮT DỄ HIỂU)", None),
        ("Phần 1: Nghiên cứu Thị trường, Đối tượng Mục tiêu & Đổi mới Sáng tạo", os.path.join(AGENTS_DIR, "specs", "market-research.md")),
        ("Phần 2: Tài liệu Phân tích Nghiệp vụ & Sản phẩm (BA Spec)", os.path.join(AGENTS_DIR, "specs", "ba.md")),
        ("Phần 3: Bốn Tính năng Đột phá Cốt lõi (Breakthroughs B1–B4)", os.path.join(AGENTS_DIR, "specs", "breakthroughs.md")),
        ("Phần 4: Chiến lược Tích hợp AI / LLM Thực dụng", os.path.join(AGENTS_DIR, "specs", "llm-strategy.md")),
        ("Phần 5: Báo cáo Rà soát Codebase & Hiện trạng Giao diện", os.path.join(AGENTS_DIR, "codebase-audit.md")),
        ("Phần 6: Kế hoạch & Lộ trình Thực hiện", os.path.join(AGENTS_DIR, "project-timeline.md")),
        ("Phần 7: Tiêu chí Đánh giá & Phân bổ Điểm số", os.path.join(AGENTS_DIR, "scoring-rubric.md")),
        ("Phần 8: Hướng dẫn Chuẩn bị Phản biện Ban Giám khảo", os.path.join(AGENTS_DIR, "judge-review.md")),
        ("Phần 9: Ý tưởng & Quyết định Kiến trúc Dự án", os.path.join(AGENTS_DIR, "ideas.md")),
    ]

    missing = [filepath for _, filepath in files[1:] if filepath and not os.path.isfile(filepath)]
    if missing:
        raise FileNotFoundError("Thiếu tài liệu nguồn: " + ", ".join(missing))

    for idx, (title, _) in enumerate(files):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Inches(0.2)
        r_num = p.add_run(f"{idx + 1}. ")
        r_num.bold = True
        r_num.font.name = "Segoe UI"
        r_num.font.color.rgb = COLOR_SECONDARY
        r_t = p.add_run(title)
        r_t.font.name = "Segoe UI"
        r_t.font.size = Pt(10.5)

    doc.add_page_break()

    # 3. Phần Tóm tắt dễ hiểu
    build_executive_summary(doc)

    # 4. Từng file chi tiết
    for idx, (title, filepath) in enumerate(files[1:], start=1):
        assert filepath is not None
        # Section Header
        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_before = Pt(16)
        p_sec.paragraph_format.space_after = Pt(8)
        p_sec.paragraph_format.keep_with_next = True
        r_sec = p_sec.add_run(title)
        r_sec.bold = True
        r_sec.font.name = "Segoe UI"
        r_sec.font.size = Pt(16)
        r_sec.font.color.rgb = COLOR_PRIMARY
        
        with open(filepath, "r", encoding="utf-8") as f:
            content = f.read()
        
        parse_markdown_to_docx(doc, content, skip_h1=True)
        doc.add_page_break()

    output = args.output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    print(f"Export successful: {output}")

if __name__ == "__main__":
    main()
