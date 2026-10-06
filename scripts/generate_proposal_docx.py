"""Build the Vietnamese/English pitch scripts and editable slide decks.

Run: python scripts/generate_proposal_docx.py
Dependencies: python-docx, python-pptx
All personas, amounts and pilot targets are labelled as examples or hypotheses.
"""

from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt as PptPt


ROOT = Path(__file__).resolve().parents[1]
NAVY = "172B4D"
TEAL = "087E83"
INK = "17253A"
MUTED = "5D6B7D"
CREAM = "F8F5EE"
WHITE = "FFFFFF"

SOURCES = [
    ("Contest scope", "IEEE CIS Shenzhen: International Track, English, Topic A and 20 October deadline", "https://r10.ieee.org/shenzhen-cis/blog/2026/06/17/2026-shenzhen-international-fintech-competition-international-track/"),
    ("Contest theme", "Securities Times reporting organizer information: AI as a Financial Participant and common core capabilities", "https://www.stcn.com/article/detail/4168818.html"),
    ("User need", "CFPB: financial well-being includes shock absorption, goals and freedom to choose; US framework, not Vietnamese user research", "https://www.consumerfinance.gov/consumer-tools/financial-well-being/about/"),
    ("Career uncertainty", "ILO 2025: occupational AI exposure usually means transformation, not certain job loss", "https://www.ilo.org/publications/generative-ai-and-jobs-2025-update"),
    ("Competitor", "YNAB: budgeting and goal targets", "https://www.ynab.com/features/goal-tracking"),
    ("Competitor", "ProjectionLab: compare life-event and career-change scenarios", "https://projectionlab.com/blog/compare-mode-upgrades"),
]


DECKS = {
    "vi": {
        "docx": "De_An_Du_An_An_Tam_AI.docx",
        "pptx": "An_Tam_AI_Pitch_VI.pptx",
        "label": "FINTECHATHON 2026 · INTERNATIONAL TRACK · TOPIC A",
        "status": "Bản pitch làm việc · 29/09/2026 · Số liệu người dùng là tình huống minh họa, chưa phải nghiên cứu khách hàng",
        "notes_label": "Lời trình bày",
        "source_label": "Cơ sở / trạng thái",
        "appendix": "Nguồn và ranh giới bằng chứng",
        "slides": [
            {
                "title": "Trước khi nghỉ việc, biết mình có kham nổi.",
                "subtitle": "An Tâm AI giúp người trẻ thử một quyết định lớn bằng tiền, nghề và mục tiêu của chính họ.",
                "points": ["Hỏi bằng lời thường.", "Thấy cái giá của từng phương án.", "Tự duyệt bước tiếp theo."],
                "note": "Pitch một câu: An Tâm biến câu hỏi ‘Tôi có nên nghỉ việc đi học?’ thành một so sánh có thể kiểm tra: còn tiền sống bao lâu, mục tiêu nào bị lùi và cần chuẩn bị thêm gì. Người dùng vẫn quyết định.",
                "evidence": "Định vị dự thi Topic A; không tuyên bố đã có người dùng trả phí.",
            },
            {
                "title": "Linh muốn nghỉ việc để học. Khi nào cô ấy làm được?",
                "subtitle": "Một quyết định; bốn thứ phải cân cùng lúc.",
                "points": ["Học phí và chi phí sống trong thời gian không lương.", "Quỹ dự phòng và mục tiêu mua nhà.", "Tiền đầu tư cùng ngành với công việc hiện tại."],
                "note": "Linh là nhân vật minh họa, chưa phải kết quả phỏng vấn. Không hỏi cô ấy có thích thêm một bảng chi tiêu không. Hỏi: nếu bắt đầu tháng tới thì lúc tệ nhất còn bao nhiêu tháng an toàn; nếu chờ hoặc học bán thời gian, điều gì thay đổi? CFPB định nghĩa an toàn tài chính bằng khả năng chịu cú sốc và quyền tự do lựa chọn, chứ không chỉ số dư.",
                "evidence": "Tình huống giả định. CFPB là khung nhu cầu tham khảo; cần phỏng vấn người dùng Việt Nam để xác nhận mức độ đau.",
            },
            {
                "title": "Một quyết định đi qua ngân sách, mục tiêu và nghề.",
                "subtitle": "Ngân sách, kế hoạch dài hạn và nghề nghiệp thường được xem tách nhau.",
                "points": ["YNAB giúp lập ngân sách và mục tiêu.", "ProjectionLab đã so sánh các kịch bản cuộc đời.", "An Tâm tìm điều kiện để thực hiện lựa chọn khi tiền và nghề cùng gặp rủi ro."],
                "note": "Đây là định vị, không phải tuyên bố không ai có tính năng tương tự. Điểm cạnh tranh cần được chứng minh bằng trải nghiệm: từ một câu hỏi đến phương án, điều kiện để thực hiện và việc tiếp theo được phê duyệt. Không dùng ma trận dấu X vô căn cứ hay câu ‘đầu tiên trên thị trường’.",
                "evidence": "Trang tính năng chính thức của YNAB và ProjectionLab; chưa có nghiên cứu so sánh người dùng trực tiếp.",
            },
            {
                "title": "Đổi cách học, khoảng thiếu giảm từ 162 xuống 78 triệu.",
                "subtitle": "Ví dụ tính tay: 84 triệu tiền dùng được · chi 21 triệu/tháng · học phí 120 triệu · 6 tháng.",
                "points": ["Nghỉ việc 6 tháng: thiếu sơ bộ 162 triệu.", "Còn thu nhập 14 triệu/tháng: thiếu sơ bộ 78 triệu.", "An Tâm xét thêm quỹ an toàn, mục tiêu và cú sốc kép."],
                "metrics": [("Nghỉ việc 6 tháng", "162 triệu"), ("Giữ thu nhập 14 triệu/tháng", "78 triệu")],
                "note": "Có thể minh họa bằng phép tính tay: 84 triệu tiền dễ rút, chi thiết yếu 21 triệu/tháng, học 6 tháng không lương và học phí 120 triệu thì khoảng thiếu sơ bộ là 162 triệu. Nếu vẫn kiếm 14 triệu/tháng, khoảng thiếu còn 78 triệu. Đây là ví dụ giải thích, không phải kết quả engine; ứng dụng còn xét quỹ dự phòng, mục tiêu và giả định khác.",
                "evidence": "Engine hiện có: shared/engine/lifeChoice.ts; ví dụ số là giả định trình bày.",
            },
            {
                "title": "Mỗi lựa chọn cần một ‘phiếu quyết định’.",
                "subtitle": "Từ cảnh báo mơ hồ thành điều kiện có thể hành động.",
                "points": ["‘Chưa đủ’ → khoảng thiếu và tháng chuẩn bị.", "‘Có điều kiện’ → giữ quỹ an toàn, giảm chi, thêm thu nhập hoặc lùi ngày bắt đầu.", "Mỗi thay đổi được thử lại ngay trên cùng phương án."],
                "note": "‘Phiếu quyết định’ là cách trình bày một phương án gồm chi phí, quỹ còn lại, mục tiêu bị ảnh hưởng, kịch bản xấu và điều kiện để tiến tới. An Tâm có thể chấm điểm khả thi, nhưng điểm số không phải nhân vật chính. Các kịch bản là phép thử độ nhạy, không phải dự báo tương lai. Không nói chắc chắn người dùng sẽ tăng lương sau khi học.",
                "evidence": "shared/engine/lifeChoice.ts có verdict, funding.shortfall, monthsAfter và choiceConditions; kiểm thử nghiệp vụ trước khi dùng số trong demo.",
            },
            {
                "title": "Điểm mới: từ ‘chưa đủ’ đến một đường đi có thể thử.",
                "subtitle": "Ba cơ chế nối với nhau trên cùng một quyết định.",
                "points": ["Tính ngược: cần thêm bao nhiêu tiền hoặc mấy tháng chuẩn bị.", "Thử sốc kép: mất lương trong lúc tài sản cùng ngành giảm giá.", "Biến điều kiện thành việc; mọi thay đổi chờ người dùng duyệt."],
                "note": "Điểm mới là vòng lặp: đánh giá một quyết định, chỉ ra điều kiện làm kết quả đổi chiều, thử tình huống lương và tài sản xấu đi cùng lúc, rồi đề xuất việc cụ thể. ILO đo mức phơi nhiễm nghề với GenAI, không dự đoán một người sẽ mất việc; dùng dữ liệu này như giả định bất định có thể xem và sửa. Không gọi đây là công nghệ ‘đầu tiên’ hay độc quyền.",
                "evidence": "shared/engine/lifeChoice.ts, concentration.ts, humanCapital.ts và các action trong shared/agent/planner.ts; ILO 2025 chỉ tham khảo phơi nhiễm nghề.",
            },
            {
                "title": "AI hiểu và sắp việc. Code tính tiền. Người dùng duyệt.",
                "subtitle": "Người dùng thấy AI đã hiểu gì, tính gì và định làm gì.",
                "points": ["Hiểu ‘tôi muốn nghỉ việc đi học’ và chọn phép tính phù hợp.", "Engine trả kết quả cùng giả định; AI diễn đạt bằng lời dễ hiểu.", "Tạo đề xuất; người dùng xem trước, duyệt hoặc từ chối."],
                "note": "Đây là cách bám chủ đề AI as a Financial Participant: AI tham gia vào chuỗi ra quyết định có dấu vết. Trong code, Workers AI chỉ chạy khi được cấu hình; nếu chậm hoặc lỗi thì có fallback. LLM không tự tính số tiền và không tự duyệt thao tác. Trên sân khấu phải cho thấy timeline và một hành động pending thật.",
                "evidence": "worker/agents/llm.ts, runner.ts, worker/routes/agent.ts, worker/agents/executor.ts; binding AI và demo end-to-end cần kiểm tra ở môi trường nộp.",
            },
            {
                "title": "Engine và luồng duyệt đã có trong sản phẩm.",
                "subtitle": "Bằng chứng kỹ thuật mạnh nhất là đường đi từ câu hỏi tới kết quả có thể kiểm tra.",
                "points": ["Đã có mô phỏng lựa chọn, sức kiếm tiền và rủi ro dồn vào một ngành.", "Đã có giao diện, lịch sử AI và bước duyệt hành động.", "Cần xác nhận dữ liệu mẫu, cấu hình AI và demo trọn luồng trước khi nộp."],
                "note": "Không gọi code hoặc test là bằng chứng đã có khách hàng. Mở demo bằng tài khoản mẫu, dùng cùng một quyết định, rồi cho giám khảo đổi một giả định và thử phê duyệt hai lần. Nếu AI binding không sẵn sàng, trình bày rõ fallback theo luật thay vì gọi đó là LLM đang chạy.",
                "evidence": "Repo hiện có các module nêu trên; production và khả năng lặp lại demo chưa được tài liệu này xác nhận.",
            },
            {
                "title": "Hai hướng thử doanh thu: Plus và tích hợp đối tác.",
                "subtitle": "Bắt đầu trực tiếp để học nhu cầu; chỉ mở rộng qua đối tác khi kết quả được chứng minh.",
                "points": ["Free để thử; Plus 49.000đ/tháng là giá đang cấu hình, chưa xác nhận sức mua.", "Pilot với ngân hàng/fintech: tích hợp vào app; phí triển khai và sử dụng là giả thuyết.", "Kiểm tra lãi gộp: 49.000đ trừ AI, hạ tầng, thanh toán và hỗ trợ mỗi user."],
                "note": "Mã nguồn đã có gói Free/Plus/Bank và luồng nâng gói qua webhook, nhưng không suy ra đã có doanh thu. Kênh trực tiếp phải đo tỷ lệ trả phí, chi phí thu hút và biên đóng góp trên mỗi người dùng. Bán cho ngân hàng là đường dài: bắt đầu bằng pilot nhỏ với dữ liệu có quyền sử dụng, bảo mật và quy trình pháp lý; chỉ đàm phán phí triển khai + phí theo người dùng hoạt động khi đối tác thấy kết quả. Không tuyên bố giảm nợ xấu hay tăng giữ chân trước khi có số đo.",
                "evidence": "shared/plans.ts có giá Plus 49.000đ; worker/lib/plans.ts có logic thanh toán. Chưa có bằng chứng hợp đồng hay doanh thu trong hồ sơ này.",
            },
            {
                "title": "Người dùng quyết định — và giải thích được vì sao.",
                "subtitle": "Kế hoạch kiểm chứng trước khi biến pitch thành lời hứa thương mại.",
                "points": ["Phỏng vấn người đang cân nhắc học, đổi việc hoặc chăm sóc gia đình.", "Cho họ thử 2 phương án; đo hiểu giả định, quyết định và bước tiếp theo.", "Pilot đối tác sau khi demo, quyền dữ liệu và chi phí/phiên đã được xác minh."],
                "note": "Kết: An Tâm không quyết định thay người dùng. Chúng tôi giúp họ thấy cái giá của lựa chọn, rủi ro khi mọi thứ xấu đi cùng lúc và điều kiện để bước tiếp. Trước hạn nộp 20/10, bài thi cần một demo lặp lại được và một bộ kiểm tra bảo mật, không cần dự phóng thị trường không có căn cứ.",
                "evidence": "Các bước pilot là kế hoạch, không phải kết quả đã đạt. Hạn International Track theo IEEE CIS Shenzhen.",
            },
        ],
    },
    "en": {
        "docx": "An_Tam_AI_Competition_Proposal_EN.docx",
        "pptx": "An_Tam_AI_Pitch_EN.pptx",
        "label": "FINTECHATHON 2026 · INTERNATIONAL TRACK · TOPIC A",
        "status": "Working pitch · 29 September 2026 · The persona and amounts are illustrative, not user-research findings",
        "notes_label": "Speaker notes",
        "source_label": "Evidence / status",
        "appendix": "Sources and evidence boundaries",
        "slides": [
            {
                "title": "Know whether you can afford the leap before you take it.",
                "subtitle": "An Tâm AI helps young adults test a major life decision against their money, career and goals.",
                "points": ["Ask in everyday language.", "See the cost of each option.", "Approve the next step yourself."],
                "note": "One-sentence pitch: An Tâm turns ‘Can I leave work to study?’ into an inspectable comparison: how long savings last, which goals move, and what must change first. The person makes the decision.",
                "evidence": "Topic A positioning; no claim of paying customers.",
            },
            {
                "title": "Linh wants to leave work to study. When can she start?",
                "subtitle": "One decision, four things to weigh together.",
                "points": ["Tuition and living costs without a salary.", "An emergency buffer and a home goal.", "Investments exposed to the same sector as her job."],
                "note": "Linh is an illustrative persona, not an interview finding. The useful question is not whether she wants another spending chart. It is how many safe months remain if she starts next month, and what changes if she waits or studies part-time. The CFPB frames financial well-being around absorbing shocks and having freedom of choice, rather than account balances alone.",
                "evidence": "Illustrative scenario. CFPB is a framing source; Vietnamese user interviews must validate the pain.",
            },
            {
                "title": "One decision spans budgets, goals and career risk.",
                "subtitle": "Budget, long-term plans and earning power are often considered separately.",
                "points": ["YNAB helps people budget and fund goals.", "ProjectionLab already compares major life scenarios.", "An Tâm focuses on the conditions for acting when career and assets can fail together."],
                "note": "This is positioning, not a claim that no competitor has similar features. The experience to prove is a single flow from question to options, conditions for action and an approved next step. Avoid unsupported comparison grids and ‘first in the market’ language.",
                "evidence": "Official YNAB and ProjectionLab product pages; no direct comparative user study yet.",
            },
            {
                "title": "Changing the study plan cuts the cash gap from 162m to 78m.",
                "subtitle": "Hand example: VND 84m liquid · VND 21m/month essentials · VND 120m tuition · six months.",
                "points": ["Six months without pay: initial gap VND 162m.", "Keep VND 14m/month income: initial gap VND 78m.", "An Tâm adds the safety buffer, goals and a joint shock."],
                "metrics": [("Six months without pay", "VND 162m"), ("Keep VND 14m/month income", "VND 78m")],
                "note": "A hand-calculated example: VND 84m liquid, VND 21m essential monthly spending, six months without pay and VND 120m tuition produces an initial VND 162m cash gap. Keeping VND 14m monthly income reduces the gap to VND 78m. These numbers explain the question; they are not engine output. The application also considers reserves, other goals and explicit assumptions.",
                "evidence": "The life-choice engine exists at shared/engine/lifeChoice.ts; numeric example is illustrative.",
            },
            {
                "title": "Every option needs a decision receipt.",
                "subtitle": "A vague warning becomes an actionable condition.",
                "points": ["‘Not yet’ → funding gap and preparation time.", "‘Possible with changes’ → protect reserves, reduce costs, add income or delay the start.", "Test each change against the same option."],
                "note": "A ‘decision receipt’ records an option's cost, remaining buffer, affected goals, adverse scenario and condition for proceeding. An Tâm can calculate a feasibility score, but the score is not the story. Scenarios are sensitivity tests, not forecasts. Do not promise a salary increase after study.",
                "evidence": "shared/engine/lifeChoice.ts exposes verdict, funding.shortfall, monthsAfter and choiceConditions; hand-check demo figures before use.",
            },
            {
                "title": "The innovation: turn ‘not yet’ into a path people can test.",
                "subtitle": "Three mechanisms work on the same decision.",
                "points": ["Work backward: how much money or preparation time is needed?", "Test a double shock: lost income while same-sector assets fall.", "Turn the condition into a task; every change waits for approval."],
                "note": "The new product loop evaluates a choice, shows what would change the outcome, tests an income and asset shock together, and proposes a concrete task. ILO occupational exposure to GenAI is not a prediction of an individual's job loss; An Tâm treats it as an inspectable uncertainty assumption. Do not claim this is the first or only tool of its kind.",
                "evidence": "shared/engine/lifeChoice.ts, concentration.ts, humanCapital.ts and actions in shared/agent/planner.ts; ILO 2025 informs exposure assumptions only.",
            },
            {
                "title": "AI interprets. Code calculates. The person approves.",
                "subtitle": "The user can see what AI understood, calculated and proposed.",
                "points": ["Understand ‘I want to leave work to study’ and select the right calculation.", "The engine returns a result and assumptions; AI explains them plainly.", "Create a proposal; the user previews, approves or rejects it."],
                "note": "This is the connection to AI as a Financial Participant: AI participates in a traceable financial workflow. In the code, Workers AI runs only when configured; errors and timeouts fall back to rules. The LLM does not calculate the amount or approve writes. The competition demo must show the timeline and a real pending action.",
                "evidence": "worker/agents/llm.ts and runner.ts; worker/routes/agent.ts; worker/agents/executor.ts. Check the AI binding and end-to-end demo in the submission environment.",
            },
            {
                "title": "The engine and approval flow already exist.",
                "subtitle": "The strongest proof is a path from question to an inspectable outcome.",
                "points": ["Decision, earning-power and sector-overlap calculations are built.", "The interface, AI history and approval step are built.", "Demo data, AI configuration and the full journey still need checks."],
                "note": "Code or tests are not evidence of customers. Open a reproducible demo account, keep one decision throughout, let the judges change an assumption and try double approval. If the AI binding is unavailable, show and label the rule-based fallback rather than implying the LLM ran.",
                "evidence": "Current repository modules listed here; this document does not verify production or demo repeatability.",
            },
            {
                "title": "Two revenue hypotheses: Plus and partner integration.",
                "subtitle": "Learn demand directly first; expand through partners only after outcomes are proven.",
                "points": ["Free invites trial; Plus is configured at VND 49,000/month, with demand untested.", "Bank/fintech pilot: planning inside their app; integration and usage fees are hypotheses.", "Test contribution: VND 49,000 minus AI, cloud, payment and support per paid user."],
                "note": "The repository contains Free/Plus/Bank plans and payment-upgrade logic, but that is not revenue evidence. The direct channel must measure paid conversion, acquisition cost and contribution margin per user. Bank distribution takes time: start with a small pilot using authorized data, security and legal review; negotiate setup and active-user fees only after measurable partner outcomes. Do not promise lower defaults or higher retention without partner data.",
                "evidence": "shared/plans.ts configures the Plus price; worker/lib/plans.ts contains payment logic. No contract or revenue is evidenced here.",
            },
            {
                "title": "We win when users can decide — and explain why.",
                "subtitle": "Validation before a commercial promise.",
                "points": ["Interview people facing study, career-change or family-care decisions.", "Let them compare two paths; test understanding, decision and next action.", "Pilot with a partner after the demo, data rights and session cost are verified."],
                "note": "Close: An Tâm does not decide for people. We show the cost of a choice, what happens when shocks compound, and the condition for moving forward. Before the 20 October deadline, the entry needs a repeatable demo and security checks, not unsupported market projections.",
                "evidence": "Pilot steps are plans, not completed results. International Track deadline: IEEE CIS Shenzhen.",
            },
        ],
    },
}


def doc_paragraph(doc, value, size=11, bold=False, color=INK, before=0, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(value)
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    return p


def build_docx(deck, path):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    doc.core_properties.title = "An Tâm AI — FinTechathon 2026 pitch script"
    doc.core_properties.author = "An Tâm Team"
    for i, slide in enumerate(deck["slides"], 1):
        if i > 1:
            doc.add_page_break()
        doc_paragraph(doc, f"{i:02d} / 10  ·  {deck['label']}", 9, True, TEAL, after=22)
        doc_paragraph(doc, slide["title"], 22, True, NAVY, after=12)
        doc_paragraph(doc, slide["subtitle"], 13, False, MUTED, after=20)
        for point in slide["points"]:
            p = doc_paragraph(doc, "•  " + point, 13, False, INK, after=13)
            p.paragraph_format.left_indent = Cm(0.35)
        doc_paragraph(doc, deck["notes_label"], 10, True, TEAL, before=24)
        doc_paragraph(doc, slide["note"], 11, after=13)
        doc_paragraph(doc, deck["source_label"], 10, True, TEAL, before=12)
        doc_paragraph(doc, slide["evidence"], 10, color=MUTED)
    doc.add_page_break()
    doc_paragraph(doc, deck["appendix"], 19, True, NAVY, after=13)
    doc_paragraph(doc, deck["status"], 10, color=MUTED, after=15)
    for label, description, url in SOURCES:
        doc_paragraph(doc, label + " — " + description, 10, True, INK, after=3)
        doc_paragraph(doc, url, 9, color=TEAL, after=9)
    doc_paragraph(doc, "Code evidence: shared/engine/lifeChoice.ts, humanCapital.ts, concentration.ts; worker/agents/llm.ts, runner.ts, executor.ts; shared/plans.ts. Recheck the final commit and running demo before submission.", 9, color=MUTED, before=16)
    doc.save(path)


def ppt_text(slide, x, y, w, h, value, size, color, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    p = frame.paragraphs[0]
    p.alignment = align
    p.text = value
    p.font.name = "Aptos"
    p.font.size = PptPt(size)
    p.font.bold = bold
    p.font.color.rgb = PptColor.from_string(color)
    return shape


def build_pptx(deck, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for i, item in enumerate(deck["slides"], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = PptColor.from_string(CREAM)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.13), Inches(7.5))
        bar.fill.solid()
        bar.fill.fore_color.rgb = PptColor.from_string(TEAL)
        bar.line.fill.background()
        ppt_text(slide, 0.68, 0.40, 11.9, 0.32, deck["label"], 10, TEAL, True)
        if i == 1:
            ppt_text(slide, 0.68, 1.17, 11.9, 2.10, item["title"], 38, NAVY, True)
            ppt_text(slide, 0.72, 3.39, 11.5, 1.10, item["subtitle"], 23, INK)
            for j, point in enumerate(item["points"]):
                ppt_text(slide, 0.76, 5.36 + j * 0.40, 11.4, 0.33, point, 15, TEAL, True)
        else:
            ppt_text(slide, 0.68, 1.02, 11.9, 1.46, item["title"], 29, NAVY, True)
            ppt_text(slide, 0.72, 2.52, 11.5, 0.82, item["subtitle"], 17, MUTED)
            if i in {3, 4, 5, 6, 7, 9}:
                for j, point in enumerate(item["points"]):
                    x = 0.70 + j * 4.13
                    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(3.62), Inches(3.82), Inches(2.12))
                    card.fill.solid()
                    card.fill.fore_color.rgb = PptColor.from_string(WHITE)
                    card.line.color.rgb = PptColor.from_string("DBE3E9")
                    ppt_text(slide, x + 0.22, 3.91, 3.30, 0.38, f"0{j + 1}", 12, TEAL, True)
                    if i == 4 and j < 2:
                        label, value = item["metrics"][j]
                        ppt_text(slide, x + 0.22, 4.34, 3.33, 0.52, label, 16, INK)
                        ppt_text(slide, x + 0.22, 4.89, 3.33, 0.52, value, 29, TEAL, True)
                    else:
                        ppt_text(slide, x + 0.22, 4.40, 3.33, 1.10, point, 17, INK, True)
            else:
                for j, point in enumerate(item["points"]):
                    y = 3.57 + j * 0.88
                    bullet = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.73), Inches(y + 0.13), Inches(0.10), Inches(0.10))
                    bullet.fill.solid()
                    bullet.fill.fore_color.rgb = PptColor.from_string(TEAL)
                    bullet.line.fill.background()
                    ppt_text(slide, 1.00, y, 11.3, 0.75, point, 19, INK)
        ppt_text(slide, 0.72, 7.09, 10.8, 0.22, deck["status"], 8, MUTED)
        ppt_text(slide, 12.08, 7.02, 0.55, 0.30, f"{i:02d}", 12, TEAL, True, PP_ALIGN.RIGHT)
        references = "\n".join(f"{name}: {url}" for name, _, url in SOURCES) if i == len(deck["slides"]) else ""
        slide.notes_slide.notes_text_frame.text = (
            f"{deck['notes_label']}: {item['note']}\n"
            f"{deck['source_label']}: {item['evidence']}\n{references}"
        )
    prs.save(path)


def main():
    for language, deck in DECKS.items():
        for key, builder in (("docx", build_docx), ("pptx", build_pptx)):
            path = ROOT / deck[key]
            builder(deck, path)
            mirror = ROOT / ".agents" / "submission" / path.name
            mirror.parent.mkdir(parents=True, exist_ok=True)
            mirror.write_bytes(path.read_bytes())
            print(f"{language} {key}: {path.name}")


if __name__ == "__main__":
    main()
